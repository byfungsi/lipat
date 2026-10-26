"""Independent python-pptx charts, plus an SVG extension with a safe raster fallback."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from zipfile import ZipFile, ZIP_DEFLATED

deck = Presentation()
for kind in [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE, XL_CHART_TYPE.AREA, XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT]:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['North & East', '日本語', 'South']
    data.add_series('Independent revenue', (12, 24, 36))
    chart = slide.shapes.add_chart(kind, Inches(1), Inches(1), Inches(8), Inches(4), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = 'Independent chart'
slide = deck.slides.add_slide(deck.slide_layouts[6])
data = XyChartData()
series = data.add_series('Unsupported scatter')
series.add_data_point(1, 2)
series.add_data_point(2, 4)
slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(8), Inches(4), data)
deck.save('tests/fixtures/python-charts.pptx')

# Use a generated Lipat artifact only for the extension/fallback mutation test.
with ZipFile('artifacts/charts-svg/column.pptx') as source:
    parts = {name: source.read(name) for name in source.namelist()}
parts['ppt/media/image1.svg'] = b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 120 80"><foreignObject><div xmlns="http://www.w3.org/1999/xhtml">Unsupported SVG profile</div></foreignObject></svg>'
with ZipFile('tests/fixtures/svg-raster-fallback.pptx', 'w', ZIP_DEFLATED) as target:
    for name, data in parts.items():
        target.writestr(name, data)
print('Generated independent chart and SVG raster-fallback fixtures.')
