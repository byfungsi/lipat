"""Generate an independent PPTX fixture using python-pptx, never Lipat's writer."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from io import BytesIO
from pathlib import Path

prs=Presentation()
prs.slide_width=Inches(13.333333)
prs.slide_height=Inches(7.5)
prs.core_properties.title='Independent fixture'
s=prs.slides.add_slide(prs.slide_layouts[6])
box=s.shapes.add_textbox(Inches(1),Inches(.7),Inches(10),Inches(1))
p=box.text_frame.paragraphs[0]
r=p.add_run();r.text='Hello ';r.font.size=Pt(32);r.font.bold=True;r.font.color.rgb=RGBColor(0x5B,0x5B,0xD6)
r=p.add_run();r.text='世界 & Lipat';r.font.size=Pt(32);r.font.italic=True
shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1),Inches(2),Inches(3),Inches(1));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(0xEE,0xEC,0xFF);shape.text='A real shape'
table=s.shapes.add_table(3,2,Inches(5),Inches(2),Inches(6),Inches(2)).table
for row,values in enumerate([['Metric','Value'],['Revenue','42'],['Team','8']]):
 for col,text in enumerate(values):table.cell(row,col).text=text
image=Image.new('RGB',(120,60),(91,91,214));buf=BytesIO();image.save(buf,format='PNG');buf.seek(0)
s.shapes.add_picture(buf,Inches(1),Inches(4),Inches(2),Inches(1))
s.notes_slide.notes_text_frame.text='These notes come from python-pptx.'
group=s.shapes.add_group_shape();group.shapes.add_shape(MSO_SHAPE.OVAL,Inches(9),Inches(4.5),Inches(1),Inches(1))
s2=prs.slides.add_slide(prs.slide_layouts[1]);s2.shapes.title.text='Inherited title geometry';s2.placeholders[1].text='Body from a layout placeholder'
Path('tests/fixtures').mkdir(parents=True,exist_ok=True)
prs.save('tests/fixtures/python-pptx.pptx')
