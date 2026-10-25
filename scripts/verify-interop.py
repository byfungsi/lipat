"""Independent consumer checks for Lipat's PPTX writer."""
import sys
from pptx import Presentation
p=Presentation(sys.argv[1])
assert len(p.slides)==2
texts='\n'.join(shape.text for slide in p.slides for shape in slide.shapes if shape.has_text_frame)
assert 'Hello 世界 & Lipat' in texts
assert 'Inherited title geometry' in texts
assert p.slides[0].notes_slide.notes_text_frame.text=='These notes come from python-pptx.'
tables=[shape.table for slide in p.slides for shape in slide.shapes if shape.has_table]
assert tables[0].cell(1,0).text=='Revenue'
assert tables[0].cell(1,1).text=='42'
images=[shape for slide in p.slides for shape in slide.shapes if shape.shape_type==13]
assert images and images[0].image.content_type=='image/png'
print('Independent python-pptx consumer verified slides, Unicode, inherited layout, notes, table cells and embedded image.')
