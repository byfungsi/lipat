"""Independently inspect template workbooks/photos and chart-fixture SVG fallback pixels."""
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from PIL import Image
from pptx import Presentation

root = Path('examples/react-demo/public/templates')
catalog = json.loads((root / 'catalog.json').read_text())
assert len(catalog) == 30
reports = []
for info in catalog:
    folder = root / info['id']
    source = json.loads((folder / 'deck.lipat.json').read_text())
    expected = [node['chart'] for slide in source['slides'] for node in slide['content']['children'] if node['type'] == 'chart']
    deck = Presentation(folder / 'deck.pptx')
    assert len(deck.slides) == 6
    charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
    assert len(charts) == len(expected), info['id']
    with ZipFile(folder / 'deck.pptx') as package:
        for chart, data in zip(charts, expected):
            assert [series.values for series in chart.series] == [tuple(series['values']) for series in data['series']], info['id']
            assert [category.label for category in chart.plots[0].categories] == data['categories'], info['id']
            stem = Path(str(chart.part.partname)).stem
            with ZipFile(BytesIO(package.read(f'ppt/embeddings/{stem}.xlsx'))) as workbook:
                sheet = ET.fromstring(workbook.read('xl/worksheets/sheet1.xml'))
                ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                cells = {cell.attrib['r']:cell for cell in sheet.findall('.//s:c',ns)}
                for i, series in enumerate(data['series']):
                    col = chr(66+i) # Template examples have at most two series.
                    assert cells[f'{col}1'].find('s:is/s:t',ns).text == series['name']
                    for j, value in enumerate(series['values']):
                        assert float(cells[f'{col}{j+2}'].find('s:v',ns).text) == value
        photographs = [name for name in package.namelist() if name.endswith(('.jpg', '.jpeg'))]
        assert photographs, info['id']
        for name in photographs:
            image = Image.open(BytesIO(package.read(name))).convert('RGB')
            assert max(image.size) <= 1280 and min(image.size) >= 700 and image.getbbox() is not None, (info['id'], name)
        for name in package.namelist():
            if name.endswith('.rels'):
                assert 'TargetMode="External"' not in package.read(name).decode(), (info['id'],name)
    reports.append({'id':info['id'],'slides':6,'charts':len(charts),'photographs':len(photographs)})
Path('artifacts/template-review/interop.json').write_text(json.dumps(reports,indent=2))
print('Independent python-pptx: 30 decks / 180 slides, chart values/categories, XLSX cells, embedded photograph pixels and offline relationships passed.')

# Templates now use photography; keep independent SVG fallback checks on the
# dedicated chart/vector fixtures instead of silently losing that coverage.
for kind in ['column', 'bar', 'line', 'area', 'pie', 'doughnut']:
    with ZipFile(Path('artifacts/charts-svg') / f'{kind}.pptx') as package:
        vectors = [name for name in package.namelist() if name.endswith('.svg')]
        assert vectors, kind
        for name in vectors:
            fallback = name[:-4] + '-fallback.png'
            image = Image.open(BytesIO(package.read(fallback))).convert('RGBA')
            assert max(image.size) <= 2048 and image.getbbox() is not None, (kind, name)
print('Six chart/SVG fixtures retain nonempty compatibility fallback pixels.')
